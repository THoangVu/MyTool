"""Tests for excel_exporter module."""

import unittest
from pathlib import Path
import openpyxl

from src.tax_invoice_downloader.models import (
    InvoiceHeader,
    InvoiceLineItem,
    DownloadHistoryRecord,
    SummaryStats,
)
from src.tax_invoice_downloader.excel_exporter import export_invoices_to_excel


class TestExcelExporter(unittest.TestCase):

    def setUp(self):
        self.test_output = Path(__file__).resolve().parent / "test_output.xlsx"
        if self.test_output.exists():
            self.test_output.unlink()

    def tearDown(self):
        if self.test_output.exists():
            self.test_output.unlink()

    def test_export_four_sheets(self):
        """Verify that export generates an Excel file with all 4 required sheets and correct contents."""
        item1 = InvoiceLineItem(
            khoa_hoa_don="3400401523_C26TNT_00005286_20261009",
            mst_nguoi_ban="3400401523",
            kh_hoa_don="C26TNT",
            sh_hoa_don="00005286",
            ngay_lap="2026-10-09",
            stt=1,
            ma_hang="SP01",
            ten_hang="Sản phẩm A",
            dvt="Cái",
            so_luong=5.0,
            don_gia=30555.56,
            thanh_tien=152778.0,
            thue_suat="8%",
            tien_thue=12222.0,
            file_xml_nguon="3400401523_C26TNT_00005286_20261009.xml",
        )
        inv = InvoiceHeader(
            mst_nguoi_ban="3400401523",
            ten_nguoi_ban="CÔNG TY TNHH NAM THẠNH",
            mst_nguoi_mua="060176009807",
            ten_nguoi_mua="HỘ KINH DOANH THỊNH ĐẠT",
            kh_mau_so="1",
            kh_hoa_don="C26TNT",
            sh_hoa_don="00005286",
            ngay_lap="2026-10-09",
            tien_chua_thue=204630.0,
            tien_thue=16370.0,
            tong_tien=221000.0,
            file_xml="3400401523_C26TNT_00005286_20261009.xml",
            trang_thai_trich_xuat="THANH_CONG",
            items=[item1],
        )
        history = [
            DownloadHistoryRecord(
                khoa_hoa_don="3400401523_C26TNT_00005286_20261009",
                trang_thai="XML_EXTRACTED",
                so_lan_thu=0,
                ten_zip_tam="invoice.zip",
                ten_xml_dau_ra="3400401523_C26TNT_00005286_20261009.xml",
            )
        ]
        stats = SummaryStats(
            so_dong_tim_thay=1,
            so_da_tai=1,
            so_xml_thanh_cong=1,
            so_loi=0,
            so_bo_qua=0,
        )

        output_path = export_invoices_to_excel([inv], history, stats, self.test_output)
        self.assertTrue(output_path.exists())

        wb = openpyxl.load_workbook(output_path)
        sheet_names = wb.sheetnames
        self.assertIn("HoaDon", sheet_names)
        self.assertIn("ChiTietSanPham", sheet_names)
        self.assertIn("LichSuTai", sheet_names)
        self.assertIn("ThongKe", sheet_names)

        # Validate HoaDon content
        ws_hd = wb["HoaDon"]
        self.assertEqual(ws_hd.cell(2, 1).value, "3400401523")
        self.assertEqual(ws_hd.cell(2, 7).value, "00005286")  # Leading zeros preserved
        self.assertEqual(float(ws_hd.cell(2, 13).value), 221000.0)

        # Validate ChiTietSanPham content
        ws_sp = wb["ChiTietSanPham"]
        self.assertEqual(ws_sp.cell(2, 7).value, "SP01")
        self.assertEqual(ws_sp.cell(2, 8).value, "Sản phẩm A")

        # Validate ThongKe content
        ws_tk = wb["ThongKe"]
        self.assertEqual(ws_tk.cell(2, 2).value, 1)


if __name__ == "__main__":
    unittest.main()
