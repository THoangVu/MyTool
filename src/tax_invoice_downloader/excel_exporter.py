"""Excel workbook exporter for tax invoices, line items, download history, and stats."""

import os
from pathlib import Path
from typing import List, Optional, Union, Dict, Any
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import InvoiceHeader, InvoiceLineItem, DownloadHistoryRecord, SummaryStats
from .xml_parser import parse_invoice_xml
from .logging_config import setup_logger, ManifestDB

logger = setup_logger("excel_exporter")
BASE_DIR = Path(__file__).resolve().parent.parent.parent
OUTPUT_DIR = BASE_DIR / "output"
DEFAULT_EXCEL_PATH = OUTPUT_DIR / "tax_invoices.xlsx"


def _apply_header_style(ws, cols_count: int):
    """Applies clean, professional corporate styling to table headers."""
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="1F4E79"),
        bottom=Side(style="medium", color="0D233A"),
    )
    for col_idx in range(1, cols_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28


def _apply_data_styles(ws, start_row: int, end_row: int, cols_count: int, num_cols: set = None, text_cols: set = None):
    """Applies alternating row fills, thin borders, and appropriate number/text formats."""
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )
    data_font = Font(name="Segoe UI", size=10)
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    num_cols = num_cols or set()
    text_cols = text_cols or set()

    for row_idx in range(start_row, end_row + 1):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, cols_count + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = alt_fill

            if col_idx in num_cols:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in text_cols:
                cell.number_format = "@"
                cell.alignment = Alignment(horizontal="center" if "Mã" in str(ws.cell(1, col_idx).value or "") or "Số" in str(ws.cell(1, col_idx).value or "") or "Ngày" in str(ws.cell(1, col_idx).value or "") else "left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")


def _auto_fit_columns(ws, max_len: int = 50):
    """Adjusts column widths dynamically according to content."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_width = 0
        for cell in col:
            val_str = str(cell.value or "")
            if cell.number_format == "#,##0" and isinstance(cell.value, (int, float)):
                val_str = f"{cell.value:,.0f}"
            max_width = max(max_width, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_width + 4, 12), max_len)


def export_invoices_to_excel(
    invoices: List[InvoiceHeader],
    history: Optional[List[DownloadHistoryRecord]] = None,
    stats: Optional[SummaryStats] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> Path:
    """Exports invoice headers, line items, download history, and summary stats to Excel.

    Creates workbook with 4 sheets:
    1. HoaDon
    2. ChiTietSanPham
    3. LichSuTai
    4. ThongKe
    """
    target_path = Path(output_path) if output_path else DEFAULT_EXCEL_PATH
    target_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # ----------------------------------------------------
    # Sheet 1: HoaDon
    # ----------------------------------------------------
    ws_hd = wb.create_sheet(title="HoaDon")
    headers_hd = [
        "Mã số thuế người bán",
        "Tên người bán",
        "Mã số thuế người mua",
        "Tên người mua",
        "Ký hiệu mẫu số",
        "Ký hiệu hóa đơn",
        "Số hóa đơn",
        "Ngày lập",
        "Tiền trước thuế",
        "Tiền thuế",
        "Chiết khấu",
        "Tiền phí",
        "Tổng thanh toán",
        "Loại tiền",
        "Trạng thái",
        "Tên file XML",
        "Trạng thái trích xuất",
    ]
    ws_hd.append(headers_hd)

    for inv in invoices:
        ws_hd.append([
            str(inv.mst_nguoi_ban or ""),
            inv.ten_nguoi_ban or "",
            str(inv.mst_nguoi_mua or ""),
            inv.ten_nguoi_mua or "",
            str(inv.kh_mau_so or ""),
            str(inv.kh_hoa_don or ""),
            str(inv.sh_hoa_don or ""),  # Keep leading zeros
            str(inv.ngay_lap or ""),
            inv.tien_chua_thue if inv.tien_chua_thue is not None else "",
            inv.tien_thue if inv.tien_thue is not None else "",
            inv.chiet_khau if inv.chiet_khau is not None else "",
            inv.phi if inv.phi is not None else "",
            inv.tong_tien if inv.tong_tien is not None else "",
            inv.loai_tien or "VND",
            inv.trang_thai or "",
            inv.file_xml or "",
            inv.trang_thai_trich_xuat or "THANH_CONG",
        ])

    _apply_header_style(ws_hd, len(headers_hd))
    if len(invoices) > 0:
        # Columns 9, 10, 11, 12, 13 are numbers
        _apply_data_styles(ws_hd, 2, len(invoices) + 1, len(headers_hd), num_cols={9, 10, 11, 12, 13}, text_cols={1, 3, 5, 6, 7, 8})
    _auto_fit_columns(ws_hd)

    # ----------------------------------------------------
    # Sheet 2: ChiTietSanPham
    # ----------------------------------------------------
    ws_sp = wb.create_sheet(title="ChiTietSanPham")
    headers_sp = [
        "Khóa hóa đơn",
        "MST người bán",
        "Ký hiệu hóa đơn",
        "Số hóa đơn",
        "Ngày lập",
        "STT",
        "Mã hàng",
        "Tên hàng",
        "Đơn vị tính",
        "Số lượng",
        "Đơn giá",
        "Thành tiền",
        "Thuế suất",
        "Tiền thuế dòng hàng",
        "Tên file XML nguồn",
    ]
    ws_sp.append(headers_sp)

    row_count_sp = 0
    for inv in invoices:
        for item in inv.items:
            row_count_sp += 1
            ws_sp.append([
                item.khoa_hoa_don or inv.invoice_key.normalized_key(),
                str(item.mst_nguoi_ban or inv.mst_nguoi_ban or ""),
                str(item.kh_hoa_don or inv.kh_hoa_don or ""),
                str(item.sh_hoa_don or inv.sh_hoa_don or ""),
                str(item.ngay_lap or inv.ngay_lap or ""),
                item.stt or "",
                str(item.ma_hang or ""),
                item.ten_hang or "",
                item.dvt or "",
                item.so_luong if item.so_luong is not None else "",
                item.don_gia if item.don_gia is not None else "",
                item.thanh_tien if item.thanh_tien is not None else "",
                str(item.thue_suat or ""),
                item.tien_thue if item.tien_thue is not None else "",
                item.file_xml_nguon or inv.file_xml or "",
            ])

    _apply_header_style(ws_sp, len(headers_sp))
    if row_count_sp > 0:
        # Col 10: SLuong, Col 11: DGia, Col 12: ThTien, Col 14: TienThue
        _apply_data_styles(ws_sp, 2, row_count_sp + 1, len(headers_sp), num_cols={10, 11, 12, 14}, text_cols={1, 2, 3, 4, 5, 6, 7, 13})
    _auto_fit_columns(ws_sp)

    # ----------------------------------------------------
    # Sheet 3: LichSuTai
    # ----------------------------------------------------
    ws_ls = wb.create_sheet(title="LichSuTai")
    headers_ls = [
        "Thời điểm chạy",
        "Khóa hóa đơn",
        "Trạng thái",
        "Số lần thử",
        "Tên ZIP tạm",
        "Tên XML đầu ra",
        "Thông báo lỗi",
    ]
    ws_ls.append(headers_ls)

    history_records = history or []
    for h in history_records:
        ws_ls.append([
            h.thoi_diem,
            h.khoa_hoa_don,
            h.trang_thai,
            h.so_lan_thu,
            h.ten_zip_tam,
            h.ten_xml_dau_ra,
            h.thong_bao_loi,
        ])

    _apply_header_style(ws_ls, len(headers_ls))
    if len(history_records) > 0:
        _apply_data_styles(ws_ls, 2, len(history_records) + 1, len(headers_ls), num_cols={4}, text_cols={1, 2, 3})
    _auto_fit_columns(ws_ls)

    # ----------------------------------------------------
    # Sheet 4: ThongKe
    # ----------------------------------------------------
    ws_tk = wb.create_sheet(title="ThongKe")
    headers_tk = ["Chỉ số", "Giá trị"]
    ws_tk.append(headers_tk)

    st = stats or SummaryStats()
    stats_data = [
        ("Số dòng hóa đơn tìm thấy", st.so_dong_tim_thay or len(invoices)),
        ("Số hóa đơn đã tải", st.so_da_tai or len(invoices)),
        ("Số XML trích xuất thành công", st.so_xml_thanh_cong or len(invoices)),
        ("Số hóa đơn lỗi", st.so_loi),
        ("Số hóa đơn bị bỏ qua do đã tồn tại", st.so_bo_qua),
    ]
    for metric, val in stats_data:
        ws_tk.append([metric, val])

    _apply_header_style(ws_tk, len(headers_tk))
    _apply_data_styles(ws_tk, 2, len(stats_data) + 1, len(headers_tk), num_cols={2})
    _auto_fit_columns(ws_tk)

    wb.save(target_path)
    logger.info(f"Xuất file Excel thành công: {target_path}")
    return target_path


def export_from_allinvoice_folder(
    xml_dir: Optional[Union[str, Path]] = None,
    output_path: Optional[Union[str, Path]] = None,
    manifest_db: Optional[ManifestDB] = None,
) -> Path:
    """Scans all XML files in ALLINVOICE/ directory, parses them, and exports to Excel.

    Uses manifest_db for download history and accurate run stats if provided.
    """
    target_xml_dir = Path(xml_dir) if xml_dir else (BASE_DIR / "ALLINVOICE")
    invoices: List[InvoiceHeader] = []

    if target_xml_dir.exists():
        for xml_file in sorted(target_xml_dir.glob("*.xml")):
            try:
                inv = parse_invoice_xml(xml_file, source_filename=xml_file.name)
                invoices.append(inv)
            except Exception as e:
                logger.warning(f"Bỏ qua file XML không hợp lệ {xml_file.name}: {e}")

    history_records: List[DownloadHistoryRecord] = []
    stats = SummaryStats()

    if manifest_db:
        hist_dicts = manifest_db.get_all_history()
        for hd in hist_dicts:
            history_records.append(DownloadHistoryRecord(
                thoi_diem=hd.get("thoi_diem", ""),
                khoa_hoa_don=hd.get("khoa_hoa_don", ""),
                trang_thai=hd.get("trang_thai", ""),
                so_lan_thu=hd.get("so_lan_thu", 0),
                ten_zip_tam=hd.get("ten_zip_tam", ""),
                ten_xml_dau_ra=hd.get("ten_xml_dau_ra", ""),
                thong_bao_loi=hd.get("thong_bao_loi", ""),
            ))
        stats = manifest_db.get_summary_stats()
    else:
        stats.so_dong_tim_thay = len(invoices)
        stats.so_da_tai = len(invoices)
        stats.so_xml_thanh_cong = len(invoices)

    return export_invoices_to_excel(
        invoices=invoices,
        history=history_records,
        stats=stats,
        output_path=output_path,
    )
